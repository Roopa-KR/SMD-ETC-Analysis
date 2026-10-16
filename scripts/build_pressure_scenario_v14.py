"""v14: synthesize a realistic evening-pressure scenario in the 10-truck demo workbook.

Why: in the v11-v13 demo every truck finished well inside its deadline, so no truck was ever At
Risk at the 18:30 starting snapshot. v14 adds a documented, explainable cause of pressure:

  * Two late add-on OTC (WT2) orders for each of two trucks, SH0908-1034 and SH0908-1069
    (orders 69001-69004: 19 lines for 1034 and 13 for 1069, released 18:20 and 18:49). Both trucks get a synthetic
    carrier appointment of 20:30, so pick cutoff = 20:00 and safety-buffer deadline = 19:50. The
    customers are existing customers of those routes; the lines copy item/location/quantity from
    existing WT2 lines. Every new row is labelled SYNTHETIC v14.
  * WT2 picking after 18:15 is re-simulated with the SAME operators, their recorded login
    sessions and each operator's measured picking pace, working the open lines in priority order
    (earliest pick cutoff first, then release time). Lines picked before 18:15 are untouched.
    Nothing is forced: statuses come from the engine's normal calculation.
  * The cause of the late finish is explicit labour data: lhall's OTC session ends at her
    RECORDED logout 19:13, so WT2 drops from two pickers to one just before the cutoff. mmathis
    works documented overtime in OTC until 20:25 (recorded logout 19:15). Even so, one line of
    each truck is still unpicked when both trucks depart at 20:30 if no action is taken.
  * No cross-training is added; recommendations use only the recorded experience table.

Idempotent: always starts from the v13 baseline copy (INFO/historical/ten_truck_v13_baseline.xlsx)
and writes data/workbook.xlsx. Deterministic: no random numbers.

    python scripts/build_pressure_scenario_v14.py [--baseline ...] [--output ...] [--report]
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
from collections import defaultdict
from pathlib import Path

import openpyxl.cell._writer as _cell_writer
from openpyxl import load_workbook
from openpyxl.compat import strings as _strings

_safe_string = _strings.safe_string


def _exact_string(value):
    """openpyxl writes floats with 16 significant digits, which turns 16:45:00 into 16:44:59.9999996.
    Write the shortest exact representation (repr) so every untouched time and rate round-trips."""
    if isinstance(value, float) and value == value and value not in (float('inf'), float('-inf')):
        return repr(value)
    return _safe_string(value)


_cell_writer.safe_string = _exact_string

ROOT = Path(__file__).resolve().parents[1]
BASELINE = ROOT / 'INFO/historical/ten_truck_v13_baseline.xlsx'
OUTPUT = ROOT / 'data/workbook.xlsx'
DAY = 46273                                   # 2026-09-08 as an Excel serial day
LABEL = 'SYNTHETIC v14 PRESSURE SCENARIO (late add-on orders; re-simulated picking)'
SCENARIO = 'V14_LATE_WAVE'
RESIM_FROM = '18:15'                          # picks after this time in RESIM_AREAS are re-simulated
RESIM_AREAS = (2,)
LOADING_MINUTES = 30

# Late add-on orders: (new order id, truck, work type, release time, lines, existing order of the
# same customer on that truck whose customer id/name is reused)
WAVES = [
    ('69001', 'SH0908-1034', 2, '18:20', 15, '61040'),     # TRUSTEE DRUGS INC (SS)
    ('69002', 'SH0908-1069', 2, '18:20', 10, '60990'),     # METTER PHARMACY (SS)
    ('69003', 'SH0908-1034', 2, '18:49', 4, '61256'),      # GREDEL PHARMACY LLC
    ('69004', 'SH0908-1069', 2, '18:49', 3, '61056'),      # VILLAGE WALK PHARMACY (SS)
]
# Synthetic carrier appointments: both trucks leave at 20:30 (pick cutoff 20:00, buffer deadline 19:50).
DEPARTURES = {'SH0908-1034': '20:30', 'SH0908-1069': '20:30'}

# Who picks in each re-simulated area: operator -> (seconds per line, available windows).
# Seconds per line = each operator's measured picking pace in the v13 replay (lhall ~27, mmathis
# ~22 lines per picking hour). lhall's window ends at her recorded logout (19:13); mmathis' window
# includes the documented overtime.
CREWS = {
    2: {'lhall': (135, [('17:23', '19:13')]),
        'mmathis': (165, [('15:33', '20:25')])},
}
# Labor-session edits that make the windows above true in LM005S1 (the data the app reads).
SESSION_CHANGES = {
    ('mmathis', 2): dict(end='20:25'),                       # overtime: recorded logout was 19:15
}
EXTRA_EXPERIENCE: list[tuple[str, int]] = []
JITTER = (0.86, 1.12, 0.94, 1.18, 0.90, 1.02, 1.08, 0.92, 1.05, 0.93)   # mean 1.0, deterministic
# Units per line of the late add-on orders (copied cyclically; typical OTC quantities).
LATE_QTY = (4, 2, 6, 1, 3, 8, 2, 5, 3, 4)


def effort(qty) -> float:
    """Relative time to pick one line (kept at 1.0: each operator's measured cadence)."""
    return 1.0


def serial(hhmm: str) -> float:
    h, m = map(int, hhmm.split(':'))
    return DAY + (h * 60 + m) / 1440


def clock(v: float) -> str:
    return (dt.datetime(1899, 12, 30) + dt.timedelta(days=v)).strftime('%H:%M:%S')


def table(ws, header_row):
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    header = list(rows[0])
    data = [list(r) for r in rows[1:] if r and r[0] is not None]
    return header, data


def write_table(ws, header_row, rows):
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    for r in rows:
        ws.append(r)


def build(baseline: Path, output: Path, report: bool = False, overrides: dict | None = None):
    """``overrides`` (development tuning only) may replace WAVES, CREWS or SESSION_CHANGES."""
    global WAVES, CREWS, SESSION_CHANGES
    if overrides:
        WAVES = overrides.get('WAVES', WAVES)
        CREWS = overrides.get('CREWS', CREWS)
        SESSION_CHANGES = overrides.get('SESSION_CHANGES', SESSION_CHANGES)
    wb = load_workbook(baseline)
    snaps = {}
    cfg = wb['CONFIG']
    for r in range(14, 32):
        n, t = cfg.cell(r, 1).value, cfg.cell(r, 2).value
        if n is not None:
            snaps[int(n)] = t if isinstance(t, (int, float)) else (t - dt.datetime(1899, 12, 30)).total_seconds() / 86400
    t0 = serial(RESIM_FROM)

    sh_h, ships = table(wb['OB_SHIPMENT'], 5)
    S = {c: i for i, c in enumerate(sh_h)}
    for r in ships:
        if r[S['SHIPMENTID']] in DEPARTURES:
            when = serial(DEPARTURES[r[S['SHIPMENTID']]])
            r[S['ORIGINAL_DEPARTURE']] = r[S['DEPARTURE_TIME']] = when
            r[S['DEPARTURE_CHANGED_AT']] = None
            r[S['PICK_CUTOFF_TIME (final)']] = when - LOADING_MINUTES / 1440
    cutoff = {r[S['SHIPMENTID']]: r[S['DEPARTURE_TIME']] - LOADING_MINUTES / 1440 for r in ships}
    or_h, orders = table(wb['OB_ORDER'], 5)
    O = {c: i for i, c in enumerate(or_h)}
    lm_h, lmap = table(wb['OB_LINE_MAP'], 5)
    L = {c: i for i, c in enumerate(lm_h)}
    tc_h, tc = table(wb['TI102C'], 2)
    C = {c: i for i, c in enumerate(tc_h)}
    ti_h, ti = table(wb['TI102'], 2)
    T = {c: i for i, c in enumerate(ti_h)}
    lab_h, lab = table(wb['LM005S1'], 2)
    B = {c: i for i, c in enumerate(lab_h)}

    order_ship = {str(r[O['WHORDERID']]): r[O['SHIPMENTID']] for r in orders}
    line_ship = {r[L['TI102ID']]: order_ship[str(r[L['WHORDERID']])] for r in lmap}
    tc_by_id = {r[C['TI102ID']]: r for r in tc}
    ti_template = {}
    for r in ti:
        ti_template.setdefault(r[T['TI102ID']], r)

    # ---------- 1. new lines (copy item, location and quantity from existing lines of that area) ----------
    next_id = max(r[C['TI102ID']] for r in tc) + 1000
    next_x3 = max(r[C['TI102X3ID']] for r in tc) + 1000
    next_assign = 6900001
    donors = {wt: [r for r in tc if r[C['WORKTYPE']] == wt] for wt in RESIM_AREAS}
    new_lines, new_orders, new_map = [], [], []
    stop_no = defaultdict(int)
    for r in orders:
        stop_no[r[O['SHIPMENTID']]] = max(stop_no[r[O['SHIPMENTID']]], int(r[O['STOPID']] or 0))
    by_order = {str(r[O['WHORDERID']]): r for r in orders}
    for oid, ship, wt, rel, n, same_customer in WAVES:
        cust, cname = by_order[same_customer][O['CUSTOMER']], by_order[same_customer][O['CUSTOMER_NAME']]
        assert by_order[same_customer][O['SHIPMENTID']] == ship
        release = serial(rel)
        stop_no[ship] += 1
        route = int(ship.split('-')[1])
        new_orders.append([oid, ship, ship, None, route, stop_no[ship], release, 1, n, str(wt), LABEL, SCENARIO,
                           cust, cname])
        for item in range(1, n + 1):
            donor = donors[wt][(item * 7 + len(new_lines)) % len(donors[wt])]
            row = list(donor)
            row[C['TI102ID']], row[C['TI102X3ID']], row[C['ASSIGNMENTID']] = next_id, next_x3, next_assign
            row[C['CREATEDATETIME']] = release
            qty = LATE_QTY[(len(new_lines)) % len(LATE_QTY)]
            row[C['ORDERQTY']] = row[C['TOTALQTY']] = row[C['PROCESSEDQTY']] = qty
            row[C['OPENQTY']] = row[C['SHORTQTY']] = row[C['CUTQTY']] = 0
            row[C['DATA_TYPE']], row[C['SCENARIO_ID']] = LABEL, SCENARIO
            new_lines.append(row)
            new_map.append([next_id, next_x3, next_assign, wt, oid, item, ship, route, stop_no[ship], LABEL,
                            cust, None, cust])
            line_ship[next_id] = ship
            next_id += 1
            next_x3 += 1
        next_assign += 1

    # ---------- 2. re-simulate WT2 picking after 18:15 in priority order ----------
    queue = defaultdict(list)
    for r in tc:
        if r[C['WORKTYPE']] in RESIM_AREAS and r[C['MOVED_TO_C_AT']] > t0:
            queue[r[C['WORKTYPE']]].append(r)
    for r in new_lines:
        queue[r[C['WORKTYPE']]].append(r)
    picked, unpicked = {}, []
    for wt, lines in queue.items():
        pending = sorted(lines, key=lambda r: (cutoff[line_ship[r[C['TI102ID']]]], r[C['CREATEDATETIME']],
                                               r[C['TI102ID']]))
        crew = {}
        for name, (seconds, windows) in CREWS[wt].items():
            prior = [x[C['MOVED_TO_C_AT']] for x in tc if x[C['RESOURCENAME']] == name and x[C['WORKTYPE']] == wt
                     and x[C['MOVED_TO_C_AT']] <= t0]
            spans = [(serial(a), serial(b)) for a, b in windows]
            crew[name] = dict(seconds=seconds, spans=spans, free=max([t0] + prior), k=0)

        def window(op, at):
            """Start of the next available time >= at and the end of that window (None = finished)."""
            for a, b in op['spans']:
                if at < b:
                    return max(at, a), b
            return None, None

        while pending:
            options = []
            for name, op in crew.items():
                start, end = window(op, op['free'])
                if start is not None:
                    options.append((start, name, end))
            if not options:
                break
            start, name, end = min(options)
            op = crew[name]
            ready = [r for r in pending if r[C['CREATEDATETIME']] + 60 / 86400 <= start]
            if not ready:
                op['free'] = min(r[C['CREATEDATETIME']] + 60 / 86400 for r in pending)
                continue
            line = ready[0]
            done = start + op['seconds'] * effort(line[C['ORDERQTY']]) * JITTER[op['k'] % len(JITTER)] / 86400
            if done > end:              # cannot finish before this window closes: try the next window
                op['free'] = end
                continue
            op['k'] += 1
            op['free'] = done
            picked[line[C['TI102ID']]] = (round(done * 86400) / 86400, name)
            pending.remove(line)
        unpicked += pending

    # ---------- 3. write TI102C ----------
    out_tc = []
    for r in tc + new_lines:
        lid = r[C['TI102ID']]
        if r[C['WORKTYPE']] in RESIM_AREAS and (lid in picked or (r in new_lines) or r[C['MOVED_TO_C_AT']] > t0):
            if lid not in picked:
                continue                # never picked in the scenario: no completion record
            row = list(r)
            row[C['MOVED_TO_C_AT']], row[C['RESOURCENAME']] = picked[lid]
            row[C['CHANGEDATETIME']] = row[C['MOVED_TO_C_AT']]
            if r not in new_lines and abs(r[C['MOVED_TO_C_AT']] - row[C['MOVED_TO_C_AT']]) > 1e-9:
                row[C['DATA_TYPE']] = LABEL
            out_tc.append(row)
        else:
            out_tc.append(list(r))
    out_tc.sort(key=lambda r: (r[C['MOVED_TO_C_AT']], r[C['TI102ID']]))

    # ---------- 4. write TI102 snapshot rows for every re-simulated or new line ----------
    affected = {r[C['TI102ID']] for rows in queue.values() for r in rows}
    out_ti = [list(r) for r in ti if not (r[T['TI102ID']] in affected and r[T['SNAPSHOT_NO']] >= min(
        n for n, t in snaps.items() if t >= t0 - 1e-9))]
    by_id_template = dict(ti_template)
    for r in [x for rows in queue.values() for x in rows]:
        if r[C['TI102ID']] in by_id_template:
            continue                    # lines without any snapshot row so far get one built from TI102C
        row = [None] * len(ti_h)
        for col in ti_h:
            if col in C:
                row[T[col]] = r[C[col]]
        by_id_template[r[C['TI102ID']]] = row
    tc_rows = {r[C['TI102ID']]: r for r in tc + new_lines}
    for lid in sorted(affected):
        src = tc_rows[lid]
        created = src[C['CREATEDATETIME']]
        done = picked.get(lid, (None, None))[0]
        for n, t in sorted(snaps.items()):
            if t < t0 - 1e-9 or created > t + 1e-9 or (done is not None and done <= t + 1e-9):
                continue
            row = list(by_id_template[lid])
            row[T['SNAPSHOT_NO']], row[T['SNAPSHOT_TIME']] = n, t
            row[T['CURRENTSTATUSID']], row[T['RESOURCENAME']] = 0, None
            qty = src[C['ORDERQTY']]
            row[T['ORDERQTY']] = row[T['TOTALQTY']] = row[T['OPENQTY']] = qty
            row[T['PROCESSEDQTY']] = row[T['SHORTQTY']] = row[T['CUTQTY']] = 0
            row[T['CREATEDATETIME']] = row[T['CHANGEDATETIME']] = created
            row[T['TRANSACTIONDATE']] = DAY
            row[T['DATA_TYPE']], row[T['SCENARIO_ID']] = LABEL, SCENARIO
            out_ti.append(row)
    out_ti.sort(key=lambda r: (r[T['SNAPSHOT_NO']], r[T['TI102ID']]))

    # ---------- 5. labour sessions (LM005S1) ----------
    out_lab = [list(r) for r in lab]
    next_lm = max(r[B['LM005ID']] for r in lab) + 1
    for (name, wt), change in SESSION_CHANGES.items():
        rows = [r for r in out_lab if r[B['RESOURCENAME']] == name and r[B['WORKTYPE']] == wt]
        main_id = max(rows, key=lambda r: r[B['TOTALTIMEINSEC']] or 0)[B['LM005ID']]
        rows = [r for r in rows if r[B['LM005ID']] == main_id]
        start = rows[0][B['STARTDATETIME']]
        live = [r for r in rows if r[B['WIPIND']] == 'Y']
        ref = max(live, key=lambda r: r[B['SNAPSHOT_TIME']])
        per_minute = ref[B['TOTALTIMEINSEC']] / ((ref[B['SNAPSHOT_TIME']] - start) * 1440)
        end = serial(change['end'])

        def stamp(row, begin, finish):
            t = row[B['SNAPSHOT_TIME']]
            active = t <= finish + 1e-9
            secs = int(round(per_minute * (min(t, finish) - begin) * 1440))
            row[B['WIPIND']] = 'Y' if active else 'N'
            row[B['ENDDATETIME']] = None if active else finish
            row[B['TOTALTIMEINSEC']], row[B['TOTALTIMEINMIN']] = secs, int(round(secs / 60))
            row[B['DATA_TYPE']], row[B['SCENARIO_ID']] = LABEL, SCENARIO

        for row in rows:
            stamp(row, start, end)
        if 'resume' in change:
            begin, finish = serial(change['resume'][0]), serial(change['resume'][1])
            for n, t in sorted(snaps.items()):
                if t + 1e-9 < begin:
                    continue
                row = list(rows[0])
                row[B['SNAPSHOT_NO']], row[B['SNAPSHOT_TIME']] = n, t
                row[B['LM005ID']], row[B['STARTDATETIME']] = next_lm, begin
                stamp(row, begin, finish)
                out_lab.append(row)
            next_lm += 1
    out_lab.sort(key=lambda r: (r[B['SNAPSHOT_NO']], r[B['LM005ID']]))

    # ---------- 6. orders, line map, shipments, experience, readme ----------
    out_orders = [list(r) for r in orders] + new_orders
    out_map = [list(r) for r in lmap] + new_map
    out_ships = [list(r) for r in ships]
    for row in out_ships:
        ship = row[S['SHIPMENTID']]
        mine = [o for o in new_orders if o[1] == ship]
        if mine:
            row[S['ORDERS']] = (row[S['ORDERS']] or 0) + len(mine)
            row[S['LINES']] = (row[S['LINES']] or 0) + sum(o[8] for o in mine)
            wts = sorted({int(w) for w in str(row[S['WORK_TYPES']]).split(',') if w.strip()} | {int(o[9]) for o in mine})
            row[S['WORK_TYPES']] = ', '.join(map(str, wts))
            row[S['OB_SCENARIO_ID']] = SCENARIO
            row[S['NOTE']] = (str(row[S['NOTE']]).split('. v14')[0] +
                              f'. v14 SYNTHETIC: carrier appointment {DEPARTURES.get(ship, "unchanged")}; late add-on orders ' +
                              ', '.join(f"{o[0]} ({o[8]} lines WT{o[9]}, released {clock(o[6])[:5]})" for o in mine) +
                              '.')
    ex_h, experience = table(wb['OPERATOR_EXPERIENCE'], 5)
    out_exp = [list(r) for r in experience]
    for name, wt in EXTRA_EXPERIENCE:
        if not any(r[0] == name and r[1] == wt for r in out_exp):
            out_exp.append([name, wt, 'Cross-trained - working knowledge', 'SIMULATED PLANNING ASSUMPTION',
                            'v14 synthetic cross-training profile for scenario analysis; not verified history.'])

    write_table(wb['TI102C'], 2, out_tc)
    write_table(wb['TI102'], 2, out_ti)
    write_table(wb['LM005S1'], 2, out_lab)
    write_table(wb['OB_ORDER'], 5, out_orders)
    write_table(wb['OB_LINE_MAP'], 5, out_map)
    write_table(wb['OB_SHIPMENT'], 5, out_ships)
    write_table(wb['OPERATOR_EXPERIENCE'], 5, out_exp)
    readme = wb['SCENARIO_README']
    rows = [list(r) for r in readme.iter_rows(values_only=True)]
    notes = {
        'Initial snapshot': '10, 18:30 on 2026-09-08; v14: SH0908-1034 and SH0908-1069 are calculated At Risk',
        'v14 pressure scenario': ('Late add-on OTC (WT2) orders 69001-69004 for SH0908-1034 and SH0908-1069, '
                                  'released 18:20 and 18:49; both trucks depart 20:30 (pick cutoff 20:00); WT2 '
                                  'picking after 18:15 re-simulated in priority order with the same operators '
                                  'and measured paces'),
        'v14 labour changes': ('lhall OTC session ends at her recorded logout 19:13; mmathis OTC overtime '
                               'until 20:25 (recorded logout 19:15)'),
        'v14 cross-training': 'none added; recommendations use only recorded operator experience',
    }
    for key, text in notes.items():
        hit = next((i for i, r in enumerate(rows) if r and r[0] == key), None)
        if hit is None:
            readme.append([key, text])
        else:
            readme.cell(hit + 1, 2).value = text
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    if report:
        for wt, lines in sorted(queue.items()):
            print(f'WT{wt}: {len(lines)} lines re-simulated after {RESIM_FROM}; '
                  f'{sum(1 for r in lines if r[C["TI102ID"]] not in picked)} never picked')
            for r in sorted(lines, key=lambda r: picked.get(r[C['TI102ID']], (9e9,))[0]):
                lid = r[C['TI102ID']]
                p = picked.get(lid)
                print(f"   {lid} {line_ship[lid]} rel {clock(r[C['CREATEDATETIME']])[:5]} "
                      f"picked {clock(p[0]) if p else '--'} {p[1] if p else ''}")
    return dict(picked=picked, unpicked=[r[C['TI102ID']] for r in unpicked], new_lines=len(new_lines))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--baseline', type=Path, default=BASELINE)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    parser.add_argument('--report', action='store_true')
    args = parser.parse_args()
    build(args.baseline, args.output, args.report)
