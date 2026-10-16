"""v11 data fix: give the 10-truck demo the same labor intensity as the real 08-Sep day.

Problem (found in the v10 investigation): the 10-truck demo kept 100% of every operator's logged
labor time (WT1 at 65%) but only 16-49% of the picked lines, so every pick rate was 2-4x too low.

Fix: each LM005S1 session's TOTALTIMEINSEC = the real-day (archived reference) value x the
work type's share of picked lines kept in the demo:

    factor(WT) = demo TI102C lines in WT / reference TI102C lines in WT   (whole day)

Session start/end times, WIPIND and operators are unchanged, so who is logged in and when stays
real; only the labor time attributed to this 10-truck workload is scaled. The old WT1 x0.65
calibration is removed. Idempotent: it always starts from the archived reference values.

Only three things in the .xlsx change: LM005S1 TOTALTIMEINSEC / TOTALTIMEINMIN / DATA_TYPE (+ its
title cell) and two SCENARIO_README rows. The edit is made directly in the sheet XML, so every
other sheet and value stays byte-for-byte identical (no re-serialisation of numbers).

Usage (development only, not run by the app):
    python scripts/rescale_demo_labor.py [--workbook data/workbook.xlsx]
            [--reference INFO/historical/historical_33_trucks_reference.xlsx]
The app rebuilds data/truck_etc.db automatically when the workbook changes.
"""
from __future__ import annotations

import argparse
import os
import re
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook

LABEL = 'SYNTHETIC LABOR REPLAY; hours scaled to share of lines kept (v11)'
TITLE = ('LM005S1 — SYNTHETIC 10-TRUCK PLANNING REPLAY (reference-derived, NOT live actual); '
         'labor time = real-day time x work type share of lines kept')
PACE_NOTE = ('ONE recorded day; synthetic ten-truck subset; labor time scaled per work type to the share of '
             'picked lines kept (v11, replaces WT1 x0.65); NOT actual monthly rate')


def _table(ws, header_row):
    rows = ws.iter_rows(min_row=header_row, values_only=True)
    header = list(next(rows))
    return header, [r for r in rows if r and r[0] is not None]


def factors(active_wb, reference_wb) -> dict[int, float]:
    def count(wb):
        header, rows = _table(wb['TI102C'], 2)
        col = header.index('WORKTYPE')
        return Counter(int(r[col]) for r in rows)
    demo, ref = count(active_wb), count(reference_wb)
    return {wt: demo.get(wt, 0) / ref[wt] for wt in sorted(ref) if ref[wt]}


def _sheet_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    book = zf.read('xl/workbook.xml').decode('utf-8')
    rels = zf.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    targets.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    sheets = re.findall(r'<(?:\w+:)?sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', book)
    return {name: targets[rid].lstrip('/') for name, rid in sheets}


CELL = re.compile(r'(<(?:\w+:)?c r="([A-Z]+[0-9]+)"[^>]*><(?:\w+:)?v>)([^<]*)(</(?:\w+:)?v>)')


def _set_cells(xml: str, values: dict[str, object]) -> str:
    """Replace the cached value of each listed cell in one pass; fail if any cell is missing."""
    seen = set()

    def swap(match):
        ref = match.group(2)
        if ref not in values:
            return match.group(0)
        seen.add(ref)
        return match.group(1) + escape(str(values[ref])) + match.group(4)

    xml = CELL.sub(swap, xml)
    missing = set(values) - seen
    if missing:
        raise ValueError(f'cells not found: {sorted(missing)[:5]}')
    return xml


def rescale(workbook: Path, reference: Path) -> dict[int, float]:
    active = load_workbook(workbook, read_only=True, data_only=True)
    ref_wb = load_workbook(reference, read_only=True, data_only=True)
    share = factors(active, ref_wb)
    rh, rrows = _table(ref_wb['LM005S1'], 2)
    ref_secs = {(r[rh.index('SNAPSHOT_NO')], r[rh.index('LM005ID')]): r[rh.index('TOTALTIMEINSEC')] for r in rrows}
    ah, arows = _table(active['LM005S1'], 2)
    letters = {name: chr(ord('A') + i) for i, name in enumerate(ah)}
    readme = [list(r) for r in active['SCENARIO_README'].iter_rows(values_only=True)]
    active.close(), ref_wb.close()

    with zipfile.ZipFile(workbook) as zf:
        paths = _sheet_paths(zf)
        lm_path, readme_path = paths['LM005S1'], paths['SCENARIO_README']
        lm_xml = zf.read(lm_path).decode('utf-8')
        rd_xml = zf.read(readme_path).decode('utf-8')
        entries = [(info, zf.read(info.filename)) for info in zf.infolist()]

    updates = {'A1': TITLE}
    for offset, row in enumerate(arows):
        excel_row = offset + 3                      # title row 1, header row 2
        key = (row[ah.index('SNAPSHOT_NO')], row[ah.index('LM005ID')])
        if key not in ref_secs:
            raise ValueError(f'LM005S1 row {key} has no matching reference row')
        seconds = int(round(ref_secs[key] * share[int(row[ah.index('WORKTYPE')])]))
        updates[f"{letters['TOTALTIMEINSEC']}{excel_row}"] = seconds
        updates[f"{letters['TOTALTIMEINMIN']}{excel_row}"] = int(round(seconds / 60))
        updates[f"{letters['DATA_TYPE']}{excel_row}"] = LABEL
    lm_xml = _set_cells(lm_xml, updates)

    pace_row = next(i + 1 for i, r in enumerate(readme) if r and r[0] == 'Operator pace')
    rd_xml = _set_cells(rd_xml, {f'B{pace_row}': PACE_NOTE})
    scaling = 'factor(WT) = demo lines / real-day lines: ' + ', '.join(f'WT{w} {f:.3f}' for w, f in share.items())
    existing = next((i + 1 for i, r in enumerate(readme) if r and r[0] == 'Labor scaling'), None)
    if existing:
        rd_xml = _set_cells(rd_xml, {f'B{existing}': scaling})
    else:
        prefix = re.search(r'<(\w+:)?sheetData', rd_xml).group(1) or ''
        n = len(readme) + 1
        new_row = (f'<{prefix}row r="{n}"><{prefix}c r="A{n}" t="str"><{prefix}v>Labor scaling</{prefix}v></{prefix}c>'
                   f'<{prefix}c r="B{n}" t="str"><{prefix}v>{escape(scaling)}</{prefix}v></{prefix}c></{prefix}row>')
        rd_xml = rd_xml.replace(f'</{prefix}sheetData>', new_row + f'</{prefix}sheetData>', 1)

    # Close the descriptor returned by mkstemp before writing/replacing. On
    # Windows, leaving it open prevents the temporary archive from being moved.
    fd, tmp_name = tempfile.mkstemp(suffix='.xlsx', dir=workbook.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        with zipfile.ZipFile(tmp, 'w') as out:
            for info, data in entries:
                if info.filename == lm_path:
                    data = lm_xml.encode('utf-8')
                elif info.filename == readme_path:
                    data = rd_xml.encode('utf-8')
                out.writestr(info, data, compress_type=info.compress_type)
        os.replace(tmp, workbook)
    finally:
        tmp.unlink(missing_ok=True)
    print(f'Rescaled {len(arows)} LM005S1 rows; factors ' + ', '.join(f'WT{w} {f:.3f}' for w, f in share.items()))
    return share


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, default=root / 'data/workbook.xlsx')
    parser.add_argument('--reference', type=Path, default=root / 'INFO/historical/historical_33_trucks_reference.xlsx')
    args = parser.parse_args()
    rescale(args.workbook, args.reference)
