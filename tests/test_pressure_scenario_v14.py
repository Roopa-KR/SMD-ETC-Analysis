"""v14: synthesized evening pressure — genuine At Risk trucks whose statuses evolve from calculations.

Scenario (scripts/build_pressure_scenario_v14.py): late add-on OTC orders for SH0908-1034 and
SH0908-1069 (both depart 20:30: pick cutoff 20:00, safety-buffer deadline 19:50); OTC picking after
18:15 re-simulated in priority order with the recorded pickers and their measured pace; lhall's
recorded 19:13 logout halves OTC capacity; mmathis works overtime until 20:25.

Without supervisor action          18:30 / 18:45 / 19:00 At Risk -> 19:15 Pick Cutoff Passed (ready after
                                    the cutoff) -> 20:00 cutoff passed with lines open -> 20:30 left short
With dclements allocated at 18:30   On Time immediately -> On Time while picking -> All Picked (scenario)
"""
from __future__ import annotations

import hashlib
import json
import re

import pytest

from calculations.excel_io import load_inputs, serial_to_datetime
from calculations.supervisor import build_state, candidates, explain_status
from calculations.dashboard_view import departed_status
from models import database as db

FIRST, SECOND = 'SH0908-1034', 'SH0908-1069'
AT_RISK, PASSED = 'AT RISK', 'PICK CUTOFF PASSED'


def state(n, moves=(), temps=(), loading=None):
    src = db.load_sources()
    return build_state(src, db.worktype_context(n), db.trucks_at(n), n, loading or {}, list(moves),
                       temporaries=list(temps))


def truck(st, ship):
    return next(t for t in st['trucks'] if t['shipment'] == ship)


def clock(serial):
    return serial_to_datetime(serial).strftime('%H:%M')


def move(name, dest, snapshot):
    src = db.load_sources()
    st = state(snapshot)
    chosen = next(c for c in candidates(src, st, dest) if c['operator'] == name)
    assert chosen['safe']
    return dict(operator=name, source=chosen['source'], dest=dest, start_snapshot=snapshot,
                destination_pace=chosen['pace'])


# ---------- 1. at least two genuinely At Risk trucks at the initial snapshot ----------
def test_two_trucks_are_at_risk_at_1830_from_the_calculation():
    st = state(10)
    risky = {t['shipment']: t for t in st['trucks'] if t['status'] in (AT_RISK, PASSED)}
    assert set(risky) == {FIRST, SECOND}
    for ship, t in risky.items():
        assert clock(t['cutoff']) == '20:00' and clock(t['deadline']) == '19:50'
        assert t['deadline'] < t['conservative_ready'] <= t['cutoff']             # 19:57: inside the buffer
        b = max((r for r in st['rows'] if r['shipment'] == ship and r['open']), key=lambda r: r['conservative_ready'])
        assert b['worktype'] == 2 and b['resources'] == 2
        # the ready time IS the calculation: now + lines ahead / cautious rate
        assert b['conservative_ready'] == pytest.approx(st['now'] + b['ahead'] / b['low_rate'] / 24)
        assert t['conservative_ready'] == b['conservative_ready']


def test_allocation_to_one_at_risk_truck_does_not_improve_the_other():
    baseline = state(10)
    allocation = move('dclements', 2, 10)
    allocation['shipment'] = FIRST
    changed = state(10, moves=[allocation])

    assert truck(baseline, FIRST)['status'] == AT_RISK
    assert truck(baseline, SECOND)['status'] == AT_RISK
    assert truck(changed, FIRST)['status'] == 'ON TIME'
    assert truck(changed, SECOND)['status'] == AT_RISK

    first_row = next(r for r in changed['rows'] if r['shipment'] == FIRST and r['worktype'] == 2)
    second_row = next(r for r in changed['rows'] if r['shipment'] == SECOND and r['worktype'] == 2)
    assert first_row['resources'] == second_row['resources'] + 1
    assert first_row['low_rate'] > second_row['low_rate']


def test_nothing_is_hardcoded_the_same_trucks_are_on_time_with_more_staff_or_time():
    relaxed = state(10, loading={FIRST: 20, SECOND: 20})        # 10 min more picking time each
    assert truck(relaxed, FIRST)['status'] == 'ON TIME' and truck(relaxed, SECOND)['status'] == 'ON TIME'


# ---------- 2. statuses evolve logically over the replay ----------
EXPECTED = {10: AT_RISK, 11: AT_RISK, 12: AT_RISK, 13: PASSED, 14: PASSED, 15: PASSED, 16: PASSED, 17: PASSED}


@pytest.mark.parametrize('ship', [FIRST, SECOND])
def test_status_timeline_without_supervisor_action(ship):
    for n, expected in EXPECTED.items():
        st = state(n)
        t = truck(st, ship)
        assert t['status'] == expected, (ship, n)
        if expected == AT_RISK:
            assert t['deadline'] < t['conservative_ready'] <= t['cutoff'] and not t['cutoff_reached']
        elif st['now'] < t['cutoff']:
            assert t['conservative_ready'] > t['cutoff'] and not t['cutoff_reached']   # projected to miss
        else:
            assert t['cutoff_reached'] and t['open_lines'] > 0                          # 20:00, 20:15: passed
    departed = next(t for t in db.trucks_at(18) if t.shipment_id == ship)
    assert departed.status == 'DEPARTED - 1 LINES LEFT BEHIND'


def test_status_changes_are_explained_by_the_data():
    st19, st1915 = state(12), state(13)
    lm = db.load_sources().lm
    lhall = lm[(lm.RESOURCENAME == 'lhall') & (lm.WORKTYPE == 2)].set_index('SNAPSHOT_NO').WIPIND
    assert lhall[12] == 'Y' and lhall[13] == 'N'                            # recorded logout at 19:13
    otc = lambda st: next(r for r in st['rows'] if r['shipment'] == FIRST and r['worktype'] == 2)
    assert otc(st19)['resources'] == 2 and otc(st1915)['resources'] == 1
    assert otc(st1915)['low_rate'] < otc(st19)['low_rate'] / 1.5
    reason = explain_status(st1915, FIRST)
    assert 'after the 20:00 pick cutoff' in reason and '1 picker' in reason and 'WT 2 OTC' in reason
    assert 'safety-buffer deadline but not after the 20:00 pick cutoff' in explain_status(st19, FIRST)
    assert 'has passed with' in explain_status(state(16), FIRST)


def test_every_status_change_has_a_changed_input():
    previous = None
    for snap in db.snapshots():
        st = state(snap.snapshot_no)
        rows = {(r['shipment'], r['worktype']): (r['open'], r['ahead'], r['resources'], r['low_rate'])
                for r in st['rows']}
        trucks = {t['shipment']: t for t in st['trucks']}
        if previous:
            for ship, t in trucks.items():
                if t['status'] != previous[1][ship]['status']:
                    before = {k: v for k, v in previous[0].items() if k[0] == ship}
                    after = {k: v for k, v in rows.items() if k[0] == ship}
                    crossed = previous[2] < t['cutoff'] <= st['now']        # the clock reached the cutoff
                    assert before != after or crossed, (ship, snap.snapshot_no)
        previous = (rows, trucks, st['now'])


# ---------- 3. consistent IDs and data across snapshots ----------
def test_truck_ids_and_line_counts_are_consistent_across_snapshots():
    data = load_inputs('data/workbook.xlsx')
    ids = set(data.shipments.SHIPMENTID)
    picked_before = {}
    for snap in db.snapshots():
        trucks = db.trucks_at(snap.snapshot_no)
        assert {t.shipment_id for t in trucks} == ids
        for t in trucks:
            assert t.lines_released == t.lines_picked + t.open_lines
            assert t.lines_picked >= picked_before.get(t.shipment_id, 0)        # picks never disappear
            picked_before[t.shipment_id] = t.lines_picked


def test_synthesized_rows_are_labelled_and_physically_consistent():
    data = load_inputs('data/workbook.xlsx')
    added = data.orders[data.orders.OB_SCENARIO_ID == 'V14_LATE_WAVE']
    assert sorted(added.WHORDERID) == ['69001', '69002', '69003', '69004']
    assert set(added.SHIPMENTID) == {FIRST, SECOND} and added.LINES.sum() == 32
    tc = data.tc
    new = tc[tc.SCENARIO_ID == 'V14_LATE_WAVE']
    assert len(new) <= 32 and new.DATA_TYPE.str.contains('SYNTHETIC v14').all()
    created = tc.set_index('TI102ID').CREATEDATETIME
    assert (tc.MOVED_TO_C_AT.values >= tc.CREATEDATETIME.values).all()      # never picked before release
    # every OTC pick after 18:15 happens while that picker is logged in to OTC
    otc = tc[(tc.WORKTYPE == 2) & (tc.MOVED_TO_C_AT > 46273 + 18.25 / 24)]
    lm = data.lm[data.lm.WORKTYPE == 2]
    for p in otc.itertuples():
        sessions = lm[(lm.RESOURCENAME == p.RESOURCENAME) & (lm.STARTDATETIME <= p.MOVED_TO_C_AT)]
        ends = sessions.ENDDATETIME.dropna()
        assert len(sessions) and (ends.empty or ends.max() >= p.MOVED_TO_C_AT - 1 / 1440), p.TI102ID
    # one operator never completes two lines at the same second
    assert not otc.duplicated(['RESOURCENAME', 'MOVED_TO_C_AT']).any()
    # open snapshot rows exist exactly between release and pick for the new lines
    snaps = dict(zip(data.snapshots['SNAPSHOT_NO'], data.snapshots['T']))
    for line in new.itertuples():
        expected = {n for n, t in snaps.items() if created[line.TI102ID] <= t < line.MOVED_TO_C_AT}
        actual = set(data.ti[data.ti.TI102ID == line.TI102ID].SNAPSHOT_NO)
        assert actual == expected, line.TI102ID


def test_generator_is_deterministic_and_reproduces_the_active_workbook(tmp_path):
    from openpyxl import load_workbook
    from scripts.build_pressure_scenario_v14 import BASELINE, build
    out = tmp_path / 'workbook.xlsx'
    build(BASELINE, out)
    a = load_workbook('data/workbook.xlsx', read_only=True, data_only=True)
    b = load_workbook(out, read_only=True, data_only=True)
    assert a.sheetnames == b.sheetnames
    for name in a.sheetnames:
        assert list(a[name].iter_rows(values_only=True)) == list(b[name].iter_rows(values_only=True)), name


# ---------- 4. reallocation recalculates, updates later snapshots, accounts for the source ----------
def test_allocation_at_1830_makes_both_trucks_safe_then_picked_by_calculation():
    m = move('dclements', 2, 10)
    now = state(10, [m])
    for ship in (FIRST, SECOND):
        t = truck(now, ship)
        assert t['status'] == 'ON TIME' and t['conservative_ready'] <= t['deadline']
        assert t['open_lines'] > 0 and not t['scenario_adjusted']     # not "All Picked" just because of the move
    for n in (11, 12, 13):
        for ship in (FIRST, SECOND):
            assert truck(state(n, [m]), ship)['status'] == 'ON TIME', (n, ship)
    at1930 = state(14, [m])
    assert truck(at1930, FIRST)['status'] == 'ALL PICKED' and truck(at1930, FIRST)['history_open_lines'] > 0
    assert 'picked by the allocated capacity' in explain_status(at1930, FIRST)
    assert truck(at1930, SECOND)['status'] == 'ON TIME'                  # still has open lines: not All Picked
    assert all(truck(state(15, [m]), s)['status'] == 'ALL PICKED' for s in (FIRST, SECOND))
    engine = {t.shipment_id: t.status for t in db.trucks_at(18)}
    final = state(18, [m])
    assert departed_status(engine[FIRST], truck(final, FIRST)) == 'DEPARTED COMPLETE'


def test_gain_is_cautious_pace_times_hours_worked_in_priority_order():
    m = move('dclements', 2, 10)
    base, moved = state(12), state(12, [m])
    gained = sum(r['scenario_picked'] for r in moved['rows'] if r['worktype'] == 2)
    slowest = min(r['empirical_low'] for r in moved['rows'] if r['worktype'] == 2)
    assert gained == int(min(m['destination_pace'], slowest) * 0.5)              # 30 min since 18:30
    # earliest deadline first: no later truck gains while an earlier one still has open OTC lines
    order = sorted((r for r in moved['rows'] if r['worktype'] == 2 and r['history_open']),
                   key=lambda r: (r['deadline'], r['shipment']))
    seen_open = False
    for r in order:
        if seen_open:
            assert r['scenario_picked'] == 0
        seen_open = seen_open or r['open'] > 0
    assert sum(r['open'] for r in moved['rows']) == sum(r['open'] for r in base['rows']) - gained


def test_source_area_loses_the_moved_operators_picks_minus_spare_capacity():
    m = move('ekim', 2, 10)                                 # ekim keeps picking Prescription in the replay
    st = state(12, [m])
    reopened = sum(r['scenario_reopened'] for r in st['rows'] if r['worktype'] == 1)
    note = next(n for n in st['scenario_notes'] if n.startswith('WT 1:'))
    picked = int(re.search(r'(\d+) line\(s\) picked by ekim', note).group(1))
    spare = int(re.search(r'spare capacity for (\d+)', note).group(1))
    assert picked > 0 and reopened == max(0, picked - spare)
    # the source area keeps every truck safe after its lines are reopened: no hidden new risk
    base = {t['shipment']: t['status'] for t in state(12)['trucks']}
    for t in st['trucks']:
        wt1 = [r for r in st['rows'] if r['shipment'] == t['shipment'] and r['worktype'] == 1 and r['open']]
        if wt1:
            assert all(r['status'] in ('ON TIME', 'ALL PICKED') for r in wt1), t['shipment']
        if base[t['shipment']] in ('ON TIME', 'ALL PICKED'):
            assert t['status'] in ('ON TIME', 'ALL PICKED'), t['shipment']


def test_allocation_that_is_not_enough_never_shows_safe():
    late = move('psam', 2, 12)               # one slow operator at 19:00 looks enough, then lhall logs out
    assert truck(state(12, [late]), FIRST)['status'] == 'ON TIME'
    later = state(14, [late])
    assert truck(later, FIRST)['status'] == PASSED and truck(later, FIRST)['open_lines'] > 0
    assert truck(later, FIRST)['scenario_adjusted']                        # psam's picks were counted
    assert 'after the 20:00 pick cutoff' in explain_status(later, FIRST)
    partial = move('dclements', 2, 13)       # at 19:15 one operator keeps the trucks At Risk, not safe
    assert {truck(state(n, [partial]), FIRST)['status'] for n in (13, 14, 15)} == {AT_RISK}


# ---------- 5. pages update immediately after confirmation ----------
@pytest.fixture
def client():
    from app import app, _SCENARIOS
    app.config['TESTING'] = True
    _SCENARIOS.clear()
    yield app.test_client()
    _SCENARIOS.clear()


def chart(html, chart_id):
    match = re.search(rf'<script id="{chart_id}-data" type="application/json">(.*?)</script>', html, re.S)
    return {s['key']: s for s in json.loads(match.group(1))['scenarios']}


def test_confirming_an_allocation_updates_dashboard_truck_planner_and_later_snapshots(client):
    url = f'/planner?snapshot=10&worktype=2&shipment={FIRST}'
    client.get(url)
    with client.session_transaction() as sess:
        token = sess['csrf_token']
    form = dict(csrf_token=token, snapshot=10, worktype=2, shipment=FIRST)
    dashboard = client.get('/?snapshot=10').get_data(as_text=True)
    assert re.search(r'At-risk trucks</span>.*?<div class="stat-number">2</div>', dashboard, re.S)
    client.post('/planner', data=dict(form, action='add', operator='dclements'))
    client.post('/planner', data=dict(form, action='allocate', scenario_version=0))
    dashboard = client.get('/?snapshot=10').get_data(as_text=True)
    assert re.search(r'At-risk trucks</span>.*?<div class="stat-number">0</div>', dashboard, re.S)
    page = client.get(f'/truck/{FIRST}?snapshot=10').get_data(as_text=True)
    rows = chart(page, 'ready-chart')
    assert rows['before']['band'] == 'yellow' and rows['confirmed']['band'] == 'green'
    assert 'Why this status: ON TIME' in page
    later = client.get(f'/truck/{FIRST}?snapshot=14').get_data(as_text=True)
    assert 'Why this status: ALL PICKED' in later and 'Scenario replay:' in later
    assert 'id="status-history"' in later and 'recorded ' in later
    csv = client.get('/export/trucks.csv?snapshot=14').get_data(as_text=True)
    assert FIRST in csv
    client.post('/planner', data=dict(form, action='reset_scenario'))
    assert 'Why this status: PICK CUTOFF PASSED' in client.get(f'/truck/{FIRST}?snapshot=14').get_data(as_text=True)


def test_status_history_table_lists_each_change_with_its_reason(client):
    page = client.get(f'/truck/{FIRST}?snapshot=16').get_data(as_text=True)
    section = page[page.index('id="status-history"'):]
    assert section.count('status change') >= 3                       # On Time -> At Risk -> Pick Cutoff Passed
    for text in ('has passed with', 'after the 20:00 pick cutoff', 'after the 19:50 safety-buffer deadline'):
        assert text in section
